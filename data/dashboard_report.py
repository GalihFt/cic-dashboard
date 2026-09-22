from __future__ import annotations

import re
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from data.dashboard_repository import DashboardRepository, format_month


REPORT_KPIS = [
    ("Muat Full (FTL - FOB)", "MUAT FULL"),
    ("Muat MT (MTL - MOB)", "MUAT MT"),
    ("Muat Full Trans (FIT - FOB)", "MUAT FULL TRANSHIPMENT"),
    ("Muat MT Trans (MIT - MOB)", "MUAT MT TRANSHIPMENT"),
    ("Bongkar Full (FOB - FXD)", "BONGKAR FULL"),
    ("Bongkar MT (MOB - MXD)", "BONGKAR MT"),
    ("Bongkar Full Trans (FOB - FIT)", "BONGKAR FULL TRANSHIPMENT"),
    ("Bongkar MT Trans (MOB - MIT)", "BONGKAR MT TRANSHIPMENT"),
]

HEADER_FILL = PatternFill("solid", fgColor="245848")
SUBHEADER_FILL = PatternFill("solid", fgColor="DDEBE5")
WHITE_FONT = Font(color="FFFFFF", bold=True)


def _sheet_name(port: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", str(port)).strip() or "Tanpa Cabang"
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        marker = f" ({suffix})"
        candidate = f"{base[:31 - len(marker)]}{marker}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _style_sheet(worksheet, header_row: int, number_columns: set[str]) -> None:
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = f"A{header_row}:{worksheet.cell(header_row, worksheet.max_column).coordinate}"
    for cell in worksheet[header_row]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = {cell.column: cell.value for cell in worksheet[header_row]}
    for column_index, title in headers.items():
        if title in number_columns:
            for row in range(header_row + 1, worksheet.max_row + 1):
                worksheet.cell(row, column_index).number_format = "#,##0"

    for column in worksheet.columns:
        values = [str(cell.value or "") for cell in column]
        width = min(max(max(map(len, values), default=0) + 2, 11), 36)
        worksheet.column_dimensions[column[0].column_letter].width = width


def _write_heading(worksheet, title: str, period: str, ports: str) -> None:
    worksheet["A1"] = title
    worksheet["A1"].font = Font(size=16, bold=True, color="173F35")
    worksheet["A2"] = f"Periode: {period}"
    worksheet["A3"] = f"Cabang: {ports}"
    for cell in (worksheet["A2"], worksheet["A3"]):
        cell.font = Font(color="546A60")


def build_dashboard_report(
    repository: DashboardRepository,
    ports: list[str],
    start_period: str,
    end_period: str,
) -> bytes:
    activities = repository.report_activity_summary(ports, start_period, end_period)
    arrivals = repository.report_vessel_arrivals(ports, start_period, end_period)
    branches = list(ports) if ports else repository.ports
    period_label = (
        format_month(start_period)
        if start_period == end_period
        else f"{format_month(start_period)} - {format_month(end_period)}"
    )
    port_label = ", ".join(branches) if ports else "Semua cabang"
    arrival_lookup = dict(zip(arrivals.get("Cabang", []), arrivals.get("Kedatangan Vessel", [])))

    overview_rows = []
    for branch in branches:
        branch_rows = activities[activities["Cabang"].eq(branch)]
        totals = branch_rows.groupby("Aktivitas")["TEU"].sum()
        record: dict[str, object] = {"Cabang": branch}
        record.update({label: int(totals.get(activity, 0)) for label, activity in REPORT_KPIS})
        record["Kedatangan Vessel"] = int(arrival_lookup.get(branch, 0))
        overview_rows.append(record)
    overview = pd.DataFrame(
        overview_rows,
        columns=["Cabang", *(label for label, _ in REPORT_KPIS), "Kedatangan Vessel"],
    )

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="Ringkasan", startrow=4, index=False)
        summary_sheet = writer.book["Ringkasan"]
        _write_heading(summary_sheet, "Ringkasan Dashboard CIC", period_label, port_label)
        _style_sheet(summary_sheet, 5, set(overview.columns) - {"Cabang"})

        used_names = {"ringkasan"}
        detail_columns = [
            "Periode",
            "Kelompok",
            "Aktivitas",
            "Transisi Status",
            "Kondisi",
            "Scope",
            "Jumlah Event",
            "Container Unik",
            "TEU",
        ]
        for branch in branches:
            sheet_name = _sheet_name(branch, used_names)
            detail = activities.loc[activities["Cabang"].eq(branch), detail_columns].copy()
            if not detail.empty:
                detail["Periode"] = detail["Periode"].map(format_month)
            detail.to_excel(writer, sheet_name=sheet_name, startrow=4, index=False)
            worksheet = writer.book[sheet_name]
            _write_heading(worksheet, f"Ringkasan Aktivitas - {branch}", period_label, branch)
            _style_sheet(
                worksheet,
                5,
                {"Jumlah Event", "Container Unik", "TEU"},
            )

    return output.getvalue()
